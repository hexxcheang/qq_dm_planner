"""Fill the DM Optical Quotation template for a customer."""

from __future__ import annotations

import calendar
import json
import re
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .pdf_render import read_image_crops, render_sheet
from .xlsx_utils import find_label, layout_item_block, replace_text

# Cell positions in the quotation template.
ATTENTION_ROWS = ("A", 8)  # company, address line 1, address line 2, person, email
Q_DATE, Q_NO, Q_VALID = "H9", "H10", "H11"
ITEMS_FIRST, ITEMS_LAST = 15, 18  # ITEMS_LAST carries the box's bottom border

# Typos in the original quotation template, corrected even in copies uploaded before the fix.
TEXT_FIXES = {"Shipping Term: Ex-work Singapore": "Shipping Term: Ex-Works Singapore"}


@dataclass
class QuoteCustomer:
    company: str = ""
    address1: str = ""
    address2: str = ""
    person: str = ""
    email: str = ""


@dataclass
class QuoteLine:
    description: str
    qty: float
    unit_price: float
    units: str = "Pcs"

    @property
    def amount(self) -> float:
        return round(self.qty * self.unit_price, 2)


@dataclass
class Quotation:
    q_date: date
    seq: int
    customer: QuoteCustomer
    lines: list[QuoteLine] = field(default_factory=list)

    @property
    def quotation_no(self) -> str:
        return quotation_no(self.q_date, self.seq)

    @property
    def valid_till(self) -> date:
        return add_one_month(self.q_date)

    @property
    def total(self) -> float:
        return round(sum(line.amount for line in self.lines), 2)


def quotation_no(d: date, seq: int) -> str:
    return f"Q{d:%y%m%d}-{seq:02d}"


def add_one_month(d: date) -> date:
    """Same day next month; the month's last day when it is shorter (31 Jan -> 28/29 Feb)."""
    year, month = (d.year + 1, 1) if d.month == 12 else (d.year, d.month + 1)
    return date(year, month, min(d.day, calendar.monthrange(year, month)[1]))


def _number(value: float) -> float | int:
    return int(value) if float(value).is_integer() else value


def build_quotation(q: Quotation, template: Path) -> Workbook:
    """The quotation template filled in for ``q`` (in memory)."""
    wb = load_workbook(template)
    ws = wb.active
    replace_text(ws, TEXT_FIXES)

    col, row = ATTENTION_ROWS
    c = q.customer
    for offset, text in enumerate([c.company, c.address1, c.address2, c.person]):
        ws[f"{col}{row + offset}"] = text or None
    email_cell = ws[f"{col}{row + 4}"]
    email_cell.value = c.email or None
    email_cell.hyperlink = f"mailto:{c.email}" if c.email else None

    ws[Q_DATE] = datetime.combine(q.q_date, datetime.min.time())
    ws[Q_NO] = q.quotation_no
    ws[Q_VALID] = datetime.combine(q.valid_till, datetime.min.time())

    rows = layout_item_block(ws, ITEMS_FIRST, ITEMS_LAST, [line.description for line in q.lines],
                             columns="AH", merges=[("B", "D")], wrap_column="B")
    for sn, (r, line) in enumerate(zip(rows, q.lines), start=1):
        ws[f"A{r}"] = sn
        ws[f"B{r}"] = line.description
        ws[f"E{r}"] = _number(line.qty)
        ws[f"F{r}"] = line.units
        ws[f"G{r}"] = round(line.unit_price, 2)
        ws[f"H{r}"] = f"=G{r}*E{r}"
        ws[f"G{r}"].number_format = ws[f"H{r}"].number_format = "#,##0.00"

    total_row = find_label(ws, "G", "Total (SGD)", start_row=ITEMS_FIRST)
    last_item = rows[-1] if rows else ITEMS_FIRST
    ws[f"H{total_row}"] = f"=SUM(H{ITEMS_FIRST}:H{last_item})"

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    return wb


def generate_quotation(q: Quotation, template: Path, out_dir: Path) -> Path:
    """Create the quotation as a PDF in ``out_dir`` and return its path."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    company = re.sub(r"[^A-Za-z0-9._-]+", "_", q.customer.company).strip("_")[:40] or "customer"
    out = out_dir / f"{q.quotation_no}_{company}.pdf"
    out.write_bytes(render_sheet(build_quotation(q, template).active, read_image_crops(template)))
    return out


class CustomerBook:
    """Customers quoted before, most recent first, so they can be picked again."""

    def __init__(self, path: Path):
        self.path = Path(path)

    def all(self) -> list[QuoteCustomer]:
        try:
            data = json.loads(self.path.read_text())
        except (OSError, json.JSONDecodeError):
            return []
        return [QuoteCustomer(**{k: str(v) for k, v in entry.items()
                                 if k in QuoteCustomer.__dataclass_fields__}) for entry in data]

    def remember(self, customer: QuoteCustomer) -> None:
        others = [c for c in self.all() if c.company.strip().lower() != customer.company.strip().lower()]
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.path.write_text(json.dumps([asdict(c) for c in [customer] + others[:49]], indent=2))
