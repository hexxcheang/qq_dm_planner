"""openpyxl helpers for growing the line-item box in the templates.

openpyxl's ``insert_rows`` moves cell values and styles but leaves merged
ranges and row heights where they were, so ``insert_rows_keep_layout`` shifts
those too. Formulas are not rewritten; callers rewrite the ones they own.
"""

from __future__ import annotations

import math
from copy import copy

from openpyxl.styles import Alignment, Border, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

LINE_HEIGHT = 15.0
MIN_FILLER_HEIGHT = 9.0


def insert_rows_keep_layout(ws: Worksheet, idx: int, amount: int) -> None:
    if amount <= 0:
        return
    below = [r for r in ws.merged_cells.ranges if r.min_row >= idx]
    for r in below:
        ws.unmerge_cells(r.coord)
    heights = {
        r: ws.row_dimensions[r].height
        for r in range(idx, ws.max_row + 1)
        if ws.row_dimensions[r].height is not None
    }
    ws.insert_rows(idx, amount)
    for r in range(idx, ws.max_row + 1):
        ws.row_dimensions[r].height = None
    for r, h in heights.items():
        ws.row_dimensions[r + amount].height = h
    for r in below:
        r.shift(0, amount)
        ws.merge_cells(r.coord)


def find_label(ws: Worksheet, column: str, label: str, start_row: int = 1) -> int:
    """Row number of the first cell in ``column`` whose text equals ``label``."""
    for row in range(start_row, ws.max_row + 1):
        value = ws[f"{column}{row}"].value
        if isinstance(value, str) and value.strip().lower() == label.lower():
            return row
    raise ValueError(f"Label {label!r} not found in column {column}")


def _text_lines(text: str, width_chars: float) -> int:
    lines = 0
    for part in str(text).splitlines() or [""]:
        lines += max(1, math.ceil(len(part) / max(width_chars, 1)))
    return lines


def layout_item_block(
    ws: Worksheet,
    first_row: int,
    last_row: int,
    descriptions: list[str],
    columns: str,
    merges: list[tuple[str, str]],
    wrap_column: str,
) -> list[int]:
    """Resize the bordered item box to hold one row per description plus a filler row.

    ``first_row``..``last_row`` is the box as drawn in the template (the last
    row carries the bottom border). Values in the box are cleared; returns the
    row numbers the caller should write the items into. The box keeps at least
    its template height, and rows grow for descriptions that wrap.
    """
    cols = [get_column_letter(i) for i in range(column_index_from_string(columns[0]),
                                                  column_index_from_string(columns[-1]) + 1)]
    template_height = sum(ws.row_dimensions[r].height or LINE_HEIGHT
                          for r in range(first_row, last_row + 1))
    first_styles = {}
    for c in cols:
        cell = ws[f"{c}{first_row}"]
        first_styles[c] = (copy(cell.border), copy(cell.font), copy(cell.alignment), cell.number_format)
    bottom_sides = {c: copy(ws[f"{c}{last_row}"].border.bottom) for c in cols}

    n_items = len(descriptions)
    base_rows = last_row - first_row + 1
    needed = n_items + 1  # items + one filler row that carries the bottom border
    if needed > base_rows:
        insert_rows_keep_layout(ws, last_row, needed - base_rows)
    last = first_row + max(needed, base_rows) - 1

    for r in list(ws.merged_cells.ranges):
        if first_row <= r.min_row <= last:
            ws.unmerge_cells(r.coord)

    for row in range(first_row, last + 1):
        for c in cols:
            border, font, alignment, number_format = first_styles[c]
            cell = ws[f"{c}{row}"]
            cell.value = None
            top = border.top if row == first_row else Side()
            bottom = bottom_sides[c] if row == last else Side()
            cell.border = Border(left=copy(border.left), right=copy(border.right),
                                 top=copy(top), bottom=copy(bottom))
            cell.font = copy(font)
            cell.alignment = Alignment(horizontal=alignment.horizontal, vertical="top",
                                       wrap_text=(c == wrap_column))
            cell.number_format = number_format
        for c1, c2 in merges:
            ws.merge_cells(f"{c1}{row}:{c2}{row}")

    wrap_start = column_index_from_string(wrap_column)
    wrap_end = next((column_index_from_string(c2) for c1, c2 in merges if c1 == wrap_column),
                    wrap_start)
    wrap_width = sum(ws.column_dimensions[get_column_letter(i)].width or 8.43
                     for i in range(wrap_start, wrap_end + 1)) * 1.1

    item_rows = list(range(first_row, first_row + n_items))
    used = 0.0
    for row, text in zip(item_rows, descriptions):
        ws.row_dimensions[row].height = LINE_HEIGHT * _text_lines(text, wrap_width)
        used += ws.row_dimensions[row].height
    for row in range(first_row + n_items, last):
        ws.row_dimensions[row].height = MIN_FILLER_HEIGHT
        used += MIN_FILLER_HEIGHT
    ws.row_dimensions[last].height = max(MIN_FILLER_HEIGHT, template_height - used)
    return item_rows
