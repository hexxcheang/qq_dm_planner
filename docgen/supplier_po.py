"""Fill the DM Optical Purchase Order template for the supplier (Hunan Chiopt)."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .pdf_render import read_image_crops, render_sheet
from .xlsx_utils import find_label, layout_item_block, replace_text

DELIVERY_DAYS = 10

# Cell positions in the PO template.
PO_NO, PO_DATE, PO_DELIVERY = "G6", "G7", "G8"
ITEMS_FIRST, ITEMS_LAST = 16, 18  # the remarks row below them closes the item box

# Typos in the original PO template, corrected even in copies uploaded before the fix.
TEXT_FIXES = {
    "PLEASE ADVICE THE DELIVERY STATUS": "PLEASE ADVISE THE DELIVERY STATUS",
    "PLEASE CONSOLE SHIPMENT WITH OTHER.": "PLEASE CONSOLIDATE SHIPMENT WITH OTHERS.",
    "SHIP VIA THIRDPARTY": "SHIP VIA THIRD-PARTY",
}


@dataclass
class POLine:
    model: str
    qty: float
    unit_price: float
    units: str = "Pcs"

    @property
    def amount(self) -> float:
        return round(self.qty * self.unit_price, 2)


@dataclass
class SupplierPO:
    po_date: date
    seq: int
    lines: list[POLine] = field(default_factory=list)

    @property
    def po_no(self) -> str:
        return po_no(self.po_date, self.seq)

    @property
    def delivery_date(self) -> date:
        return self.po_date + timedelta(days=DELIVERY_DAYS)

    @property
    def total(self) -> float:
        return round(sum(line.amount for line in self.lines), 2)


def po_no(d: date, seq: int) -> str:
    return f"PO{d:%y%m%d}-{seq:02d}"


def _number(value: float) -> float | int:
    return int(value) if float(value).is_integer() else value


def build_supplier_po(po: SupplierPO, template: Path) -> Workbook:
    """The PO template filled in for ``po`` (in memory)."""
    wb = load_workbook(template)
    ws = wb.active
    replace_text(ws, TEXT_FIXES)

    ws[PO_NO] = po.po_no
    ws[PO_DATE] = datetime.combine(po.po_date, datetime.min.time())
    ws[PO_DELIVERY] = datetime.combine(po.delivery_date, datetime.min.time())

    rows = layout_item_block(ws, ITEMS_FIRST, ITEMS_LAST, [line.model for line in po.lines],
                             columns="AH", merges=[("B", "C")], wrap_column="B")
    for sn, (row, line) in enumerate(zip(rows, po.lines), start=1):
        ws[f"A{row}"] = sn
        ws[f"B{row}"] = line.model
        ws[f"D{row}"] = _number(line.qty)
        ws[f"E{row}"] = line.units
        ws[f"G{row}"] = round(line.unit_price, 2)
        ws[f"H{row}"] = f"=D{row}*G{row}"
        ws[f"G{row}"].number_format = ws[f"H{row}"].number_format = "0.00"

    total_row = find_label(ws, "G", "Total (USD)", start_row=ITEMS_FIRST)
    last_item = rows[-1] if rows else ITEMS_FIRST
    ws[f"H{total_row}"] = f"=SUM(H{ITEMS_FIRST}:H{last_item})"

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    return wb


def generate_supplier_po(po: SupplierPO, template: Path, out_dir: Path) -> Path:
    """Create the PO as a PDF in ``out_dir`` and return its path."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"{re.sub(r'[^A-Za-z0-9._-]+', '_', po.po_no)}_Chiopt.pdf"
    out.write_bytes(render_sheet(build_supplier_po(po, template).active, read_image_crops(template)))
    return out
