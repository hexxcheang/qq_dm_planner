"""Fill the DM Optical Invoice and Delivery Order templates from a PO."""

from __future__ import annotations

import re
import shutil
import subprocess
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from .models import Customer, PurchaseOrder
from .numbering import do_no, invoice_no
from .xlsx_utils import find_label, layout_item_block

ROOT = Path(__file__).resolve().parent.parent
INVOICE_TEMPLATE = ROOT / "templates" / "DMO_Invoice_Template.xlsx"
DO_TEMPLATE = ROOT / "templates" / "DMO_DO_Template.xlsx"

# Cell positions in the templates.
INV_PO_NO, INV_DATE, INV_NO, INV_DO_NO = "C5", "H10", "H11", "H12"
INV_BILL_TO_ROWS, INV_SHIP_TO_ROWS = ("A", 12), ("D", 12)  # name, address1, address2, attention
INV_ITEMS_FIRST, INV_ITEMS_LAST = 18, 20

DO_DATE, DO_NO, DO_PO_NO = "G9", "G10", "G11"
DO_SHIP_TO_ROWS = ("A", 10)
DO_ITEMS_FIRST, DO_ITEMS_LAST = 16, 17


@dataclass
class GeneratedDocs:
    invoice_no: str
    do_no: str
    invoice_xlsx: Path
    do_xlsx: Path
    invoice_pdf: Path | None = None
    do_pdf: Path | None = None


def _number(value: float) -> float | int:
    return int(value) if float(value).is_integer() else value


def _fit_one_page(ws) -> None:
    # Long POs push the footer past the page edge; shrink to one page instead.
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1


def _write_customer(ws, anchor: tuple[str, int], customer: Customer) -> None:
    col, row = anchor
    for offset, text in enumerate(customer.lines()):
        ws[f"{col}{row + offset}"] = text or None


def build_invoice(po: PurchaseOrder, doc_date: date, seq: int, out_path: Path,
                  template: Path = INVOICE_TEMPLATE) -> Path:
    wb = load_workbook(template)
    ws = wb.active

    ws[INV_PO_NO] = po.po_number
    ws[INV_DATE] = datetime.combine(doc_date, datetime.min.time())
    ws[INV_NO] = invoice_no(doc_date, seq)
    ws[INV_DO_NO] = do_no(doc_date, seq)
    if po.customer:
        _write_customer(ws, INV_BILL_TO_ROWS, po.customer)
        _write_customer(ws, INV_SHIP_TO_ROWS, po.customer)

    rows = layout_item_block(ws, INV_ITEMS_FIRST, INV_ITEMS_LAST,
                             [i.description for i in po.items],
                             columns="AH", merges=[("B", "D")], wrap_column="B")
    for sn, (row, item) in enumerate(zip(rows, po.items), start=1):
        ws[f"A{row}"] = sn
        ws[f"B{row}"] = item.description
        ws[f"E{row}"] = _number(item.qty)
        ws[f"F{row}"] = item.units
        ws[f"G{row}"] = _number(item.unit_price)
        ws[f"H{row}"] = f"=G{row}*E{row}"

    total_row = find_label(ws, "G", "Total", start_row=INV_ITEMS_FIRST)
    due_row = find_label(ws, "G", "Amount Due", start_row=total_row)
    last_item = rows[-1] if rows else INV_ITEMS_FIRST
    ws[f"H{total_row}"] = f"=SUM(H{INV_ITEMS_FIRST}:H{last_item})"
    ws[f"H{due_row}"] = f"=H{total_row}"
    _fit_one_page(ws)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def build_delivery_order(po: PurchaseOrder, doc_date: date, seq: int, out_path: Path,
                         template: Path = DO_TEMPLATE) -> Path:
    wb = load_workbook(template)
    ws = wb.active

    ws[DO_DATE] = datetime.combine(doc_date, datetime.min.time())
    ws[DO_NO] = do_no(doc_date, seq)
    ws[DO_PO_NO] = po.po_number
    if po.customer:
        _write_customer(ws, DO_SHIP_TO_ROWS, po.customer)

    rows = layout_item_block(ws, DO_ITEMS_FIRST, DO_ITEMS_LAST,
                             [i.description for i in po.items],
                             columns="AG", merges=[("D", "E"), ("F", "G")], wrap_column="D")
    for sn, (row, item) in enumerate(zip(rows, po.items), start=1):
        ws[f"A{row}"] = sn
        ws[f"B{row}"] = _number(item.qty)
        ws[f"C{row}"] = item.units
        ws[f"D{row}"] = item.description
        ws[f"F{row}"] = "N/A"
    _fit_one_page(ws)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def _safe(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", text).strip("_") or "PO"


def find_soffice() -> str | None:
    for name in ("soffice", "libreoffice"):
        path = shutil.which(name)
        if path:
            return path
    for candidate in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                      r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
                      "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if Path(candidate).exists():
            return candidate
    return None


def to_pdf(xlsx: Path) -> Path | None:
    """Convert with LibreOffice if it is installed; returns None otherwise."""
    soffice = find_soffice()
    if not soffice:
        return None
    try:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir",
                        str(xlsx.parent), str(xlsx)],
                       check=True, capture_output=True, timeout=120)
    except (subprocess.SubprocessError, OSError):
        return None
    pdf = xlsx.with_suffix(".pdf")
    return pdf if pdf.exists() else None


def generate_documents(po: PurchaseOrder, doc_date: date, seq: int, out_dir: Path,
                       make_pdf: bool = False) -> GeneratedDocs:
    inv_no, d_no = invoice_no(doc_date, seq), do_no(doc_date, seq)
    po_part = _safe(po.po_number)
    inv_path = build_invoice(po, doc_date, seq, Path(out_dir) / f"{inv_no}_{po_part}.xlsx")
    do_path = build_delivery_order(po, doc_date, seq, Path(out_dir) / f"{d_no}_{po_part}.xlsx")
    docs = GeneratedDocs(inv_no, d_no, inv_path, do_path)
    if make_pdf:
        docs.invoice_pdf = to_pdf(inv_path)
        docs.do_pdf = to_pdf(do_path)
    return docs
