"""Read a customer Purchase Order into a ``PurchaseOrder``.

Two readers:

* ``parse_with_claude`` sends the PO (PDF, image, Excel or CSV) to Claude and
  gets back structured JSON. It copes with any PO layout, including scans.
  Needs ``ANTHROPIC_API_KEY``.
* ``parse_with_rules`` is free and works offline. It reads text-based PDFs,
  Excel and CSV files directly, and scanned PDFs and photos through OCR
  (see ``ocr.py``). It finds the PO number by its label and the item table by
  its header row (Description / Qty / Unit / Unit Price / Amount).

Either way the result is shown in the app for checking before anything is
generated.
"""

from __future__ import annotations

import base64
import bisect
import csv
import io
import os
import re
from pathlib import Path

from pydantic import BaseModel, Field

from . import ocr
from .models import LineItem, PurchaseOrder

PDF_TYPES = {".pdf"}
IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
               ".webp": "image/webp", ".gif": "image/gif"}
SHEET_TYPES = {".xlsx", ".xlsm"}
CSV_TYPES = {".csv"}
SUPPORTED_TYPES = sorted(PDF_TYPES | set(IMAGE_TYPES) | SHEET_TYPES | CSV_TYPES)

CLAUDE_MODEL = "claude-opus-5-5"


class ParseError(Exception):
    pass


# --------------------------------------------------------------------------- Claude


class _POLine(BaseModel):
    description: str = Field(description="Item description exactly as written on the PO, "
                                         "including part / model numbers")
    quantity: float
    units: str = Field(description="Unit of measure, e.g. Pcs, Set, Box. Use 'Pcs' if not stated")
    unit_price: float = Field(description="Price per unit, without currency symbol")


class _POExtract(BaseModel):
    po_number: str = Field(description="The customer's purchase order number")
    items: list[_POLine]


_PROMPT = (
    "This is a Purchase Order sent by a customer to our company, DM Optical Pte Ltd "
    "(the supplier/vendor). Extract the PO number and every ordered line item. "
    "Copy descriptions as written on the PO. Quantities and prices are plain numbers. "
    "Do not invent items; skip subtotal, GST/tax, delivery and total rows."
)


def claude_available() -> bool:
    return bool(os.environ.get("ANTHROPIC_API_KEY") or os.environ.get("ANTHROPIC_AUTH_TOKEN"))


def _content_block(data: bytes, suffix: str) -> dict:
    if suffix in PDF_TYPES:
        return {"type": "document", "source": {"type": "base64", "media_type": "application/pdf",
                                               "data": base64.standard_b64encode(data).decode()}}
    if suffix in IMAGE_TYPES:
        return {"type": "image", "source": {"type": "base64", "media_type": IMAGE_TYPES[suffix],
                                            "data": base64.standard_b64encode(data).decode()}}
    rows = _rows_from_sheet(data) if suffix in SHEET_TYPES else _rows_from_csv(data)
    text = "\n".join("\t".join(c for c in row) for row in rows)
    return {"type": "text", "text": f"Purchase Order (spreadsheet, tab-separated):\n{text}"}


def parse_with_claude(data: bytes, filename: str) -> PurchaseOrder:
    import anthropic

    suffix = Path(filename).suffix.lower()
    client = anthropic.Anthropic()
    response = client.beta.messages.parse(
        model=CLAUDE_MODEL,
        max_tokens=16000,
        betas=["server-side-fallback-2026-07-01"],
        fallbacks="default",
        messages=[{"role": "user", "content": [_content_block(data, suffix),
                                               {"type": "text", "text": _PROMPT}]}],
        output_format=_POExtract,
    )
    if response.stop_reason == "refusal" or response.parsed_output is None:
        raise ParseError("Claude could not read this PO. Please enter the details manually.")
    extract = response.parsed_output
    return PurchaseOrder(
        po_number=extract.po_number.strip(),
        items=[LineItem(description=short_description(i.description), qty=i.quantity,
                        units=(i.units or "Pcs").strip() or "Pcs", unit_price=i.unit_price)
               for i in extract.items],
    )


# --------------------------------------------------------------------------- rules

_PO_LABEL = re.compile(
    r"(?:P\.?\s?O\.?|Purchase\s+Order)\s*(?:No\.?|Number|Num\.?|#)?\s*[:#.]?\s*"
    r"([A-Z0-9][A-Z0-9\-/_.]*\d[A-Z0-9\-/_]*)",
    re.IGNORECASE,
)

_HEADER_KEYS = {
    "description": ("description", "desc", "item description", "particulars", "product",
                    "part no", "part number", "item"),
    "qty": ("qty", "quantity", "order qty", "qty ordered"),
    "units": ("uom", "unit", "units", "u/m"),
    "unit_price": ("unit price", "price", "unit cost", "rate", "u/price", "price/unit"),
    "amount": ("amount", "total", "ext price", "extended price", "line total", "net amount"),
}


def _norm(text) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip().lower().rstrip(":.")


def _to_number(text) -> float | None:
    if text is None:
        return None
    if isinstance(text, (int, float)):
        return float(text)
    cleaned = re.sub(r"[^\d.\-]", "", str(text).replace(",", ""))
    if cleaned in ("", "-", ".", "-."):
        return None
    try:
        return float(cleaned)
    except ValueError:
        return None


def _match_header(row: list[str]) -> dict[str, int] | None:
    cells = [_norm(c) for c in row]
    found: dict[str, int] = {}
    for field in ("unit_price", "amount", "qty", "units", "description"):  # most specific first
        for key in _HEADER_KEYS[field]:  # keys in order of preference
            idx = next((i for i, cell in enumerate(cells)
                        if cell and i not in found.values()
                        and (cell == key or cell.startswith(key + " ") or cell.endswith(" " + key))),
                       None)
            if idx is not None:
                found[field] = idx
                break
    if "description" in found and "qty" in found:
        return found
    return None


_TABLE_END = re.compile(r"^(sub\s*-?\s*total|grand\s+total|total|tax\s+code|gst|tax\b)", re.I)
_NOTES_START = re.compile(r"^(\*|remarks?\b|notes?\b|shipping\b|delivery\b|amendment\b|\d+\))",
                          re.I)
_PART_NO = re.compile(r"\bP\s*/\s*N\s*[:.#]?\s*([A-Z0-9][A-Z0-9\-./]*)", re.I)


def short_description(text: str) -> str:
    """The invoice/DO description for a PO line, without remarks and addresses.

    "ML5010C / MACRO LENS ... / P/N:73620230 / *CUSTOMER ... / REMARK ..."
    becomes "73620230 (ML5010C)": part number with the model code, as on the
    existing invoices. Otherwise the description lines before any remarks.
    """
    lines = [ln.strip() for ln in str(text).splitlines() if ln.strip()]
    if not lines:
        return ""
    part_no = _PART_NO.search(text)
    code = lines[0]
    if part_no and " " not in code and re.search(r"\d", code) and re.search(r"[A-Za-z]", code):
        return f"{part_no.group(1)} ({code})"
    kept = []
    for line in lines:
        if kept and _NOTES_START.match(line):
            break
        kept.append(line)
    return "\n".join(kept[:3])


def items_from_rows(rows: list[list[str]]) -> list[LineItem]:
    """Read the item lines under every item-table header in ``rows``.

    Lines with a description but no quantity continue the previous item's
    description (wrapped text, part numbers, remarks); ``short_description``
    trims those afterwards. A Total / Tax line ends the table.
    """
    items: list[LineItem] = []
    cols: dict[str, int] | None = None
    table_items = 0
    for line in rows:
        header = _match_header(line)
        if header:
            cols, table_items = header, 0
            continue
        if cols is None:
            continue
        get = lambda f: line[cols[f]] if f in cols and cols[f] < len(line) else None  # noqa: E731
        desc = str(get("description") or "").strip()
        qty = _to_number(get("qty"))
        if not any(str(c or "").strip() for c in line):
            continue
        first_text = desc or next(str(c).strip() for c in line if str(c or "").strip())
        if _TABLE_END.match(first_text):
            cols = None
            continue
        if desc and qty is not None and qty > 0:
            price = _to_number(get("unit_price"))
            amount = _to_number(get("amount"))
            if price is None and amount is not None:
                price = round(amount / qty, 4)
            units = str(get("units") or "").strip() or "Pcs"
            items.append(LineItem(description=desc, qty=qty, units=units, unit_price=price or 0.0))
            table_items += 1
        elif table_items and desc and qty is None and _to_number(get("amount")) is None:
            items[-1].description += "\n" + desc
        elif table_items:
            cols = None  # something that is not an item line: the table has ended
    for item in items:
        item.description = short_description(item.description)
    return items


def find_po_number(text: str) -> str:
    for match in _PO_LABEL.finditer(text):
        return match.group(1).strip(".")
    return ""


def _rows_from_sheet(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    rows = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            rows.append(["" if v is None else str(v) for v in row])
    return rows


def _rows_from_csv(data: bytes) -> list[list[str]]:
    text = data.decode("utf-8-sig", errors="replace")
    return [row for row in csv.reader(io.StringIO(text))]


def _pdf_tables_and_text(data: bytes) -> tuple[list[list[list[str]]], str]:
    """Candidate row sets (ruled tables first, then text-aligned columns) and the page text."""
    import pdfplumber

    ruled: list[list[str]] = []
    aligned: list[list[str]] = []
    texts: list[str] = []
    text_strategy = {"vertical_strategy": "text", "horizontal_strategy": "text"}
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            texts.append(page.extract_text() or "")
            for table in page.extract_tables():
                ruled.extend([["" if c is None else str(c) for c in r] for r in table])
            for table in page.extract_tables(text_strategy):
                aligned.extend([["" if c is None else str(c) for c in r] for r in table])
    return [ruled, aligned], "\n".join(texts)


def _rows_from_ocr(pages: list[list["ocr.Box"]]) -> tuple[list[list[str]], str]:
    """Turn OCR boxes into table rows and plain text.

    Above the item table each text line becomes a row of its pieces. Once the
    header line is found, the left edges of its column titles split every
    following line into the same columns, so right-aligned numbers still land
    under the right title.
    """
    rows: list[list[str]] = []
    text_lines: list[str] = []
    for boxes in pages:
        starts: list[float] | None = None
        for line in ocr.group_lines(boxes):
            cells = [b.text for b in line]
            text_lines.append(" ".join(cells))
            if _match_header(cells):
                starts = [b.x0 for b in line]
                rows.append(cells)
                continue
            if starts is None:
                rows.append(cells)
                continue
            row = [""] * len(starts)
            for box in line:
                col = max(0, bisect.bisect_right(starts, box.xc) - 1)
                row[col] = f"{row[col]} {box.text}".strip()
            rows.append(row)
    return rows, "\n".join(text_lines)


def _ocr_po(pages: list[list["ocr.Box"]]) -> PurchaseOrder:
    rows, text = _rows_from_ocr(pages)
    return PurchaseOrder(po_number=find_po_number(text), items=items_from_rows(rows))


_NO_OCR = ("Reading scanned POs and photos needs the free OCR package. Run "
           "start_app.bat again (or: pip install -r requirements.txt), or enter the items manually.")


def parse_with_rules(data: bytes, filename: str) -> tuple[PurchaseOrder, bool]:
    """Returns (po, used_ocr)."""
    suffix = Path(filename).suffix.lower()
    if suffix in PDF_TYPES:
        candidates, text = _pdf_tables_and_text(data)
        if not text.strip():  # scanned: the pages are only pictures
            if not ocr.available():
                raise ParseError("This PDF is a scan. " + _NO_OCR)
            return _ocr_po(ocr.ocr_pdf(data)), True
        items = next((found for rows in candidates if (found := items_from_rows(rows))), [])
        return PurchaseOrder(po_number=find_po_number(text), items=items), False
    if suffix in IMAGE_TYPES:
        if not ocr.available():
            raise ParseError(_NO_OCR)
        return _ocr_po([ocr.ocr_image_bytes(data)]), True
    if suffix in SHEET_TYPES:
        rows = _rows_from_sheet(data)
    elif suffix in CSV_TYPES:
        rows = _rows_from_csv(data)
    else:
        raise ParseError(f"Unsupported file type {suffix!r}. Use one of: {', '.join(SUPPORTED_TYPES)}")
    text = "\n".join(" ".join(r) for r in rows)
    return PurchaseOrder(po_number=find_po_number(text), items=items_from_rows(rows)), False


def parse_po(data: bytes, filename: str, use_claude: bool | None = None) -> tuple[PurchaseOrder, str]:
    """Parse with Claude when available, else with rules. Returns (po, reader_name)."""
    if use_claude is None:
        use_claude = claude_available()
    if use_claude:
        return parse_with_claude(data, filename), "Claude"
    po, used_ocr = parse_with_rules(data, filename)
    return po, "built-in, scan read with OCR" if used_ocr else "built-in"
