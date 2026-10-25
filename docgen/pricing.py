"""Supplier price lookup for POs to Hunan Chiopt.

Two price lists are used, read straight from the Excel files:

* DMO lens price list ("DMO Lens Price List_JM Buying Price"): maps
  "Product Code (DM Optical)" to "Model No. (Chiopt)".
* Chiopt price list ("Lens Price_Chiopt to DM"): gives the
  "DM Buying Price (USD)" for each "型号（Model No.）".

A code typed in the app can be either a DM product code or a Chiopt model
number. Prices are read from the values Excel last calculated and saved,
rounded to cents.

Customer quotations use the DMO list alone: the "JM Buying Price (SGD)"
divided by QUOTE_FACTOR, described by product code and series heading.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

MODEL_HEADER = re.compile(r"model\s*no", re.I)
DM_PRICE_HEADER = re.compile(r"dm\s*buying\s*price\s*\(usd\)", re.I)
PRODUCT_CODE_HEADER = re.compile(r"product\s*code\s*\(dm\s*optical\)", re.I)
CHIOPT_MODEL_HEADER = re.compile(r"model\s*no\.?\s*\(chiopt\)", re.I)
JM_PRICE_HEADER = re.compile(r"jm\s*buying\s*price\s*\(sgd\)", re.I)

QUOTE_FACTOR = 0.65  # customer quotation price = JM buying price (SGD) / 0.65


def normalize(code) -> str:
    """Codes compare without case, spaces or a trailing '.0' from Excel numbers."""
    if code is None:
        return ""
    if isinstance(code, float) and code.is_integer():
        code = int(code)
    return re.sub(r"\s+", "", str(code)).upper()


@dataclass
class ChioptPrice:
    model: str
    price_usd: float | None  # None: the model is listed but has no DM buying price
    sheet: str


@dataclass
class Resolved:
    code: str
    model: str | None
    price_usd: float | None
    matched_by: str  # "DM product code", "Chiopt model no." or "" when not found
    note: str = ""


def _text(value) -> str:
    return str(value).strip() if value is not None else ""


def read_chiopt_prices(path: Path) -> dict[str, ChioptPrice]:
    """Model no. -> DM buying price (USD), from every sheet of the Chiopt list.

    Each sheet has one or more tables whose header row starts with
    "型号（Model No.）". The "DM Buying Price (USD)" column is named in the
    sheet's first header row and applies to every table below it.
    """
    wb = load_workbook(path, data_only=True, read_only=True)
    prices: dict[str, ChioptPrice] = {}
    for ws in wb.worksheets:
        price_col: int | None = None
        in_table = False
        for row in ws.iter_rows(values_only=True):
            first = _text(row[0]) if row else ""
            if MODEL_HEADER.search(first):
                in_table = True
                for i, value in enumerate(row):
                    if DM_PRICE_HEADER.search(_text(value)):
                        price_col = i
                continue
            if not in_table or not first or len(first) > 30 or " " in first.strip():
                continue  # section titles such as ' 2/3" FA系列', notes, blanks
            price = row[price_col] if price_col is not None and price_col < len(row) else None
            price = round(float(price), 2) if isinstance(price, (int, float)) and price > 0 else None
            key = normalize(first)
            if key not in prices or (prices[key].price_usd is None and price is not None):
                prices[key] = ChioptPrice(first, price, ws.title)
    wb.close()
    return prices


@dataclass
class DmoItem:
    code: str | None  # DM product code as shown in the list; None when it has none
    model: str  # Chiopt model no.
    jm_price_sgd: float | None
    series: str  # the series heading above the item, e.g. '1/1.8" 5MP FA Series'
    sheet: str


def _code_text(value) -> str | None:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = _text(value)
    return text if text and text != "-" else None


def read_dmo_catalog(path: Path) -> list[DmoItem]:
    """Every item in the DMO lens price list, with its series heading.

    Each sheet has a header row naming "Model No. (Chiopt)" and usually
    "Product Code (DM Optical)" and "JM Buying Price (SGD)". A row whose only
    content is a text in column A is a series heading for the items below it.
    """
    wb = load_workbook(path, data_only=True, read_only=True)
    items: list[DmoItem] = []
    for ws in wb.worksheets:
        code_col = model_col = price_col = None
        series = ""
        for row in ws.iter_rows(values_only=True):
            texts = [_text(v) for v in row]
            filled = [i for i, t in enumerate(texts) if t]
            if filled == [0] and isinstance(row[0], str):
                series = texts[0]
                continue
            if any(CHIOPT_MODEL_HEADER.search(t) for t in texts):
                model_col = next(i for i, t in enumerate(texts) if CHIOPT_MODEL_HEADER.search(t))
                code_col = next((i for i, t in enumerate(texts) if PRODUCT_CODE_HEADER.search(t)),
                                None)
                price_col = next((i for i, t in enumerate(texts) if JM_PRICE_HEADER.search(t)),
                                 None)
                continue
            if model_col is None or model_col >= len(texts) or not texts[model_col]:
                continue
            code = _code_text(row[code_col]) if code_col is not None and code_col < len(row) else None
            price = row[price_col] if price_col is not None and price_col < len(row) else None
            price = float(price) if isinstance(price, (int, float)) and price > 0 else None
            items.append(DmoItem(code, texts[model_col], price, series.strip(), ws.title))
    wb.close()
    return items


def read_product_codes(path: Path) -> dict[str, str]:
    """DM product code -> Chiopt model no., from every sheet of the DMO list."""
    codes: dict[str, str] = {}
    for item in read_dmo_catalog(path):
        if item.code:
            codes.setdefault(normalize(item.code), item.model)
    return codes


@dataclass
class QuoteResolved:
    code: str  # as typed
    item: DmoItem | None
    unit_price: float | None  # JM buying price / QUOTE_FACTOR, rounded to cents
    description: str
    matched_by: str
    note: str = ""


class DmoCatalog:
    """Customer quotation prices: JM Buying Price (SGD) / QUOTE_FACTOR."""

    def __init__(self, items: list[DmoItem]):
        self.by_code: dict[str, DmoItem] = {}
        self.by_model: dict[str, DmoItem] = {}
        for item in items:
            if item.code:
                self.by_code.setdefault(normalize(item.code), item)
            key = normalize(item.model)
            # Prefer the entry that has a product code and a price.
            current = self.by_model.get(key)
            if current is None or (_rank(item) > _rank(current)):
                self.by_model[key] = item

    @classmethod
    def from_file(cls, path: Path) -> "DmoCatalog":
        return cls(read_dmo_catalog(path))

    def resolve(self, code) -> QuoteResolved:
        key, shown = normalize(code), _text(code)
        if key in self.by_code:
            item, matched_by = self.by_code[key], "DM product code"
        elif key in self.by_model:
            item, matched_by = self.by_model[key], "Chiopt model no."
        else:
            return QuoteResolved(shown, None, None, shown, "", "Not found in the DMO price list")
        notes = []
        if item.code:
            description = item.code
        else:
            description = item.model
            notes.append("no DM product code in the list – model no. used")
        if item.series:
            description = f"{description}, {item.series}"
        price = round(item.jm_price_sgd / QUOTE_FACTOR, 2) if item.jm_price_sgd else None
        if price is None:
            notes.append("no JM buying price – enter the price")
        return QuoteResolved(shown, item, price, description, matched_by, "; ".join(notes))


def _rank(item: DmoItem) -> int:
    return (item.code is not None) + (item.jm_price_sgd is not None)


def looks_like_chiopt_list(path: Path) -> bool:
    return workbook_mentions(path, DM_PRICE_HEADER)


def looks_like_dmo_list(path: Path) -> bool:
    return workbook_mentions(path, PRODUCT_CODE_HEADER)


def workbook_mentions(path: Path, pattern: re.Pattern, max_rows: int = 15) -> bool:
    try:
        wb = load_workbook(path, read_only=True)
    except Exception:  # not an Excel file openpyxl can read
        return False
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(max_row=max_rows, values_only=True):
                if any(pattern.search(_text(v)) for v in row):
                    return True
        return False
    finally:
        wb.close()


class PriceBook:
    def __init__(self, chiopt_prices: dict[str, ChioptPrice], product_codes: dict[str, str]):
        self.prices = chiopt_prices
        self.codes = product_codes

    @classmethod
    def from_files(cls, chiopt_path: Path, dmo_path: Path | None) -> "PriceBook":
        codes = read_product_codes(dmo_path) if dmo_path else {}
        return cls(read_chiopt_prices(chiopt_path), codes)

    def resolve(self, code) -> Resolved:
        key = normalize(code)
        shown = _text(code)
        if not key:
            return Resolved(shown, None, None, "", "")
        if key in self.codes:
            model, matched_by = self.codes[key], "DM product code"
        elif key in self.prices:
            model, matched_by = self.prices[key].model, "Chiopt model no."
        else:
            return Resolved(shown, None, None, "", "Not found in either price list")
        entry = self.prices.get(normalize(model))
        if entry is None:
            return Resolved(shown, model, None, matched_by,
                            f"{model} is not in the Chiopt price list – enter the price")
        if entry.price_usd is None:
            return Resolved(shown, entry.model, None, matched_by,
                            f"No DM buying price for {entry.model} – enter the price")
        return Resolved(shown, entry.model, entry.price_usd, matched_by)
