"""Draw a filled-in template worksheet straight to PDF, without Excel or LibreOffice.

The templates stay the single source of layout: column widths, row heights,
merged cells, borders, fonts, alignment and the logo are read from the
worksheet and drawn with reportlab. This covers what the DM Optical
templates use (text, numbers, dates, simple formulas, borders, one image),
not every Excel feature.

Calibri is drawn with Carlito, a free font with the same character widths,
so text lines up as it does in Excel.
"""

from __future__ import annotations

import io
import math
import re
import zipfile
from xml.etree import ElementTree
from datetime import date, datetime
from pathlib import Path

from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

FONT_DIR = Path(__file__).resolve().parent / "fonts"
REGULAR, BOLD = "Carlito", "Carlito-Bold"

EMU_PER_PT = 12700
MAX_DIGIT_PX = 7  # Calibri 11, Excel's default font
PADDING = 2.0  # pt between a cell edge and its text
LINE_SPACING = 1.22  # Calibri line height relative to font size
# Excel measures hinted screen fonts, a little narrower than print, so text
# that fits in Excel can be a few points too wide here. Where no border is in
# the way that slack is allowed; next to a border, wrapping is strict.
WRAP_SLACK = 4.0
BORDER_WIDTHS = {"hair": 0.25, "thin": 0.6, "medium": 1.2, "thick": 1.8, "double": 0.5,
                 "dashed": 0.6, "dotted": 0.6, "mediumDashed": 1.2}

# Chinese text needs a CJK font. These ship with Windows / macOS / common Linux
# desktops; the first one that loads is embedded (subset) in the PDF.
CJK_FONT_CANDIDATES = [
    (r"C:\Windows\Fonts\msyh.ttc", 0),  # Microsoft YaHei
    (r"C:\Windows\Fonts\msyh.ttf", 0),
    (r"C:\Windows\Fonts\simsun.ttc", 0),  # SimSun
    (r"C:\Windows\Fonts\simhei.ttf", 0),
    (r"C:\Windows\Fonts\Deng.ttf", 0),
    ("/System/Library/Fonts/Supplemental/Arial Unicode.ttf", 0),
    ("/Library/Fonts/Arial Unicode.ttf", 0),
    ("/System/Library/Fonts/STHeiti Light.ttc", 0),
    ("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc", 0),
    ("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc", 0),
    ("/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf", 0),
]
CJK_TEXT = re.compile(r"[\u2e80-\u9fff\uf900-\ufaff\uff00-\uffef\u3000-\u303f]")

_fonts_registered = False
_cjk_font: str | None = None


def cjk_font() -> str:
    """Name of a registered font that can draw Chinese characters."""
    global _cjk_font
    if _cjk_font is None:
        for path, index in CJK_FONT_CANDIDATES:
            if not Path(path).exists():
                continue
            try:
                pdfmetrics.registerFont(TTFont("CJK", path, subfontIndex=index))
            except Exception:  # e.g. a font with CFF outlines, which reportlab cannot embed
                continue
            _cjk_font = "CJK"
            break
        else:
            from reportlab.pdfbase.cidfonts import UnicodeCIDFont

            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))  # PDF viewers supply it
            _cjk_font = "STSong-Light"
    return _cjk_font


def _register_fonts() -> None:
    global _fonts_registered
    if not _fonts_registered:
        pdfmetrics.registerFont(TTFont(REGULAR, str(FONT_DIR / "Carlito-Regular.ttf")))
        pdfmetrics.registerFont(TTFont(BOLD, str(FONT_DIR / "Carlito-Bold.ttf")))
        _fonts_registered = True


# --------------------------------------------------------------------------- geometry


def _col_width_pt(ws: Worksheet, col: int) -> float:
    dim = ws.column_dimensions.get(get_column_letter(col))
    width = dim.width if dim is not None and dim.width else None
    if width is None:
        width = ws.sheet_format.defaultColWidth or (ws.sheet_format.baseColWidth or 8) + 0.43
    return math.trunc(width * MAX_DIGIT_PX + 0.5) * 0.75


def _row_height_pt(ws: Worksheet, row: int) -> float:
    dim = ws.row_dimensions.get(row)
    if dim is not None and dim.height is not None:
        return dim.height
    return ws.sheet_format.defaultRowHeight or 15.0


def _used_range(ws: Worksheet) -> tuple[int, int]:
    """Last row and column that hold a value or a border."""
    last_row = last_col = 1
    for row in ws.iter_rows():
        for cell in row:
            b = cell.border
            if cell.value is not None or any(s is not None and s.style for s in
                                             (b.left, b.right, b.top, b.bottom)):
                last_row = max(last_row, cell.row)
                last_col = max(last_col, cell.column)
    for rng in ws.merged_cells.ranges:
        last_row, last_col = max(last_row, rng.max_row), max(last_col, rng.max_col)
    return last_row, last_col


# --------------------------------------------------------------------------- values


_REF = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")
_SUM = re.compile(r"SUM\(([^()]*)\)", re.I)
_SAFE = re.compile(r"^[\d.+\-*/() eE]*$")


def _evaluate(ws: Worksheet, formula: str, depth: int = 0):
    """Evaluate simple formulas: cell references, + - * /, SUM(ranges)."""
    if depth > 20:
        return None
    expr = formula.lstrip("=")

    def value_of(ref: str) -> float:
        v = cell_value(ws, ws[ref], depth + 1)
        return float(v) if isinstance(v, (int, float)) else 0.0

    def sum_args(match: re.Match) -> str:
        total = 0.0
        for part in match.group(1).split(","):
            part = part.strip().replace("$", "")
            if ":" in part:
                min_col, min_row, max_col, max_row = range_boundaries(part)
                for r in range(min_row, max_row + 1):
                    for c in range(min_col, max_col + 1):
                        total += value_of(f"{get_column_letter(c)}{r}")
            elif part:
                total += value_of(part)
        return repr(total)

    expr = _SUM.sub(sum_args, expr)
    expr = _REF.sub(lambda m: repr(value_of(m.group(1) + m.group(2))), expr)
    if not _SAFE.match(expr):
        return None
    try:
        return eval(expr, {"__builtins__": {}}, {})  # only digits and arithmetic remain
    except (SyntaxError, ZeroDivisionError, TypeError):
        return None


def cell_value(ws: Worksheet, cell, depth: int = 0):
    value = cell.value
    if isinstance(value, str) and value.startswith("="):
        return _evaluate(ws, value, depth)
    return value


def format_value(value, number_format: str) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        fmt = number_format.lower()
        if "mmm" in fmt:
            year = f"{value:%Y}" if "yyyy" in fmt else f"{value:%y}"
            return f"{value.day}-{value:%b}-{year}"
        return f"{value.day}/{value.month}/{value:%Y}"
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        fmt = number_format or "General"
        if fmt == "General":
            if float(value).is_integer():
                return str(int(value))
            return f"{value:.10g}"
        decimals = sum(1 for ch in fmt.split(";")[0].split(".")[-1] if ch == "0") if "." in fmt else 0
        text = f"{abs(value):,.{decimals}f}" if "," in fmt else f"{abs(value):.{decimals}f}"
        if "$" in fmt:
            text = "$" + text
        if value < 0:
            text = f"({text})" if "(" in fmt else "-" + text
        return text
    return str(value)


# --------------------------------------------------------------------------- drawing


def _font_name(font) -> str:
    return BOLD if font is not None and font.b else REGULAR


def _font_size(font) -> float:
    return float(font.sz) if font is not None and font.sz else 11.0


def _font_rgb(font) -> tuple[float, float, float]:
    color = getattr(font, "color", None) if font is not None else None
    rgb = getattr(color, "rgb", None) if color is not None else None
    if isinstance(rgb, str) and len(rgb) in (6, 8):
        rgb = rgb[-6:]
        return tuple(int(rgb[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return (0.0, 0.0, 0.0)


def _wrap(text: str, font: str, size: float, width: float) -> list[str]:
    lines: list[str] = []
    for paragraph in text.split("\n"):
        words = paragraph.split(" ")
        current = ""
        for word in words:
            candidate = f"{current} {word}" if current else word
            if current and pdfmetrics.stringWidth(candidate, font, size) > width:
                lines.append(current)
                current = word
            else:
                current = candidate
        lines.append(current)
    return lines


class _Layout:
    def __init__(self, ws: Worksheet, last_row: int, last_col: int):
        self.col_x = [0.0]
        for c in range(1, last_col + 2):
            self.col_x.append(self.col_x[-1] + _col_width_pt(ws, c))
        self.row_y = [0.0]
        for r in range(1, last_row + 2):
            self.row_y.append(self.row_y[-1] + _row_height_pt(ws, r))
        self.width = self.col_x[last_col]
        self.height = self.row_y[last_row]

    def x(self, col: int) -> float:  # left edge of 1-based column
        return self.col_x[col - 1]

    def y(self, row: int) -> float:  # top edge of 1-based row, measured downwards
        return self.row_y[row - 1]


def read_image_crops(xlsx_path: Path) -> list[tuple[float, float, float, float] | None]:
    """Picture crops (left, top, right, bottom fractions) from the first sheet's drawing.

    openpyxl drops a picture's crop when it loads a workbook, so it is read
    from the original file. One entry per picture, in drawing order.
    """
    ns = {"xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
          "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    with zipfile.ZipFile(xlsx_path) as z:
        names = sorted(n for n in z.namelist() if re.fullmatch(r"xl/drawings/drawing\d+\.xml", n))
        if not names:
            return []
        root = ElementTree.fromstring(z.read(names[0]))
    crops = []
    for pic in root.iter(f"{{{ns['xdr']}}}pic"):
        rect = pic.find("xdr:blipFill/a:srcRect", ns)
        if rect is None:
            crops.append(None)
            continue
        crops.append(tuple(int(rect.get(k, "0")) / 100000 for k in ("l", "t", "r", "b")))
    return crops


def render_sheet(ws: Worksheet,
                 image_crops: list[tuple[float, float, float, float] | None] | None = None) -> bytes:
    """Render one worksheet as a one-page A4 PDF and return the PDF bytes."""
    _register_fonts()
    last_row, last_col = _used_range(ws)
    layout = _Layout(ws, last_row, last_col)

    page_w, page_h = A4
    margins = ws.page_margins
    left, right = margins.left * 72, margins.right * 72
    top, bottom = margins.top * 72, margins.bottom * 72
    scale = min(1.0, (page_w - left - right) / layout.width, (page_h - top - bottom) / layout.height)

    buf = io.BytesIO()
    pdf = canvas.Canvas(buf, pagesize=A4)
    pdf.translate(left, page_h - top)
    pdf.scale(scale, scale)
    # From here on: x to the right, y = -distance from the top of the sheet.

    merged_anchor = {}
    merged_inner = set()
    for rng in ws.merged_cells.ranges:
        merged_anchor[(rng.min_row, rng.min_col)] = rng
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    merged_inner.add((r, c))

    _draw_images(pdf, ws, layout, image_crops or [])
    _draw_borders(pdf, ws, layout, last_row, last_col)
    for row in ws.iter_rows(min_row=1, max_row=last_row, max_col=last_col):
        for cell in row:
            if (cell.row, cell.column) in merged_inner:
                continue
            value = cell_value(ws, cell)
            if value is None or value == "":
                continue
            rng = merged_anchor.get((cell.row, cell.column))
            max_row, max_col = (rng.max_row, rng.max_col) if rng else (cell.row, cell.column)
            box = (layout.x(cell.column), layout.y(cell.row),
                   layout.x(max_col + 1), layout.y(max_row + 1))
            right_edge = ws.cell(cell.row, max_col).border.right
            next_left = ws.cell(cell.row, max_col + 1).border.left
            bordered = bool((right_edge and right_edge.style) or (next_left and next_left.style))
            _draw_text(pdf, cell, value, box, bordered)

    pdf.showPage()
    pdf.save()
    return buf.getvalue()


def _draw_text(pdf: canvas.Canvas, cell, value, box, bordered: bool = False) -> None:
    x0, y0, x1, y1 = box
    font, size = _font_name(cell.font), _font_size(cell.font)
    text = format_value(value, cell.number_format)
    if CJK_TEXT.search(text):
        font = cjk_font()
    align = cell.alignment
    horizontal = align.horizontal or ("right" if isinstance(value, (int, float, date, datetime))
                                      and not isinstance(value, bool) else "left")
    if align.wrap_text:
        room = x1 - x0 - 2 * PADDING + (0 if bordered else WRAP_SLACK)
        lines = _wrap(text, font, size, max(1.0, room))
    else:
        lines = [text.replace("\n", " ")]

    line_h = size * LINE_SPACING
    block_h = line_h * len(lines)
    vertical = align.vertical or "bottom"
    if vertical == "top":
        block_top = y0 + 1
    elif vertical == "center":
        block_top = y0 + ((y1 - y0) - block_h) / 2
    else:
        block_top = y1 - block_h - 1

    pdf.setFont(font, size)
    pdf.setFillColorRGB(*_font_rgb(cell.font))
    for i, line in enumerate(lines):
        baseline = block_top + i * line_h + size * 0.98
        if horizontal == "right":
            pdf.drawRightString(x1 - PADDING, -baseline, line)
        elif horizontal in ("center", "centerContinuous"):
            pdf.drawCentredString((x0 + x1) / 2, -baseline, line)
        else:
            pdf.drawString(x0 + PADDING, -baseline, line)


def _draw_borders(pdf: canvas.Canvas, ws: Worksheet, layout: _Layout, last_row: int,
                  last_col: int) -> None:
    # Inside a merged range only its outer edges are drawn, as in Excel.
    merged_of = {}
    for rng in ws.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                merged_of[(r, c)] = rng
    pdf.setStrokeColorRGB(0, 0, 0)
    for row in ws.iter_rows(min_row=1, max_row=last_row, max_col=last_col):
        for cell in row:
            b = cell.border
            rng = merged_of.get((cell.row, cell.column))
            inner = set()
            if rng is not None:
                if cell.column > rng.min_col:
                    inner.add("left")
                if cell.column < rng.max_col:
                    inner.add("right")
                if cell.row > rng.min_row:
                    inner.add("top")
                if cell.row < rng.max_row:
                    inner.add("bottom")
            x0, y0 = layout.x(cell.column), layout.y(cell.row)
            x1, y1 = layout.x(cell.column + 1), layout.y(cell.row + 1)
            for side, (ax, ay, bx, by) in (("left", (x0, y0, x0, y1)), ("right", (x1, y0, x1, y1)),
                                           ("top", (x0, y0, x1, y0)), ("bottom", (x0, y1, x1, y1))):
                style = getattr(b, side).style if getattr(b, side) is not None else None
                if not style or side in inner:
                    continue
                pdf.setLineWidth(BORDER_WIDTHS.get(style, 0.6))
                pdf.setDash(2, 2) if ("dash" in style.lower() or style == "dotted") else pdf.setDash()
                pdf.line(ax, -ay, bx, -by)
    pdf.setDash()


def _draw_images(pdf: canvas.Canvas, ws: Worksheet, layout: _Layout,
                 crops: list[tuple[float, float, float, float] | None]) -> None:
    from PIL import Image

    for index, img in enumerate(getattr(ws, "_images", [])):
        anchor = img.anchor
        start = getattr(anchor, "_from", None)
        if start is None:
            continue
        x0 = layout.x(start.col + 1) + start.colOff / EMU_PER_PT
        y0 = layout.y(start.row + 1) + start.rowOff / EMU_PER_PT
        end = getattr(anchor, "to", None)
        if end is not None:
            x1 = layout.x(end.col + 1) + end.colOff / EMU_PER_PT
            y1 = layout.y(end.row + 1) + end.rowOff / EMU_PER_PT
        else:
            ext = anchor.ext
            x1, y1 = x0 + ext.width / EMU_PER_PT, y0 + ext.height / EMU_PER_PT
        picture = Image.open(io.BytesIO(img._data()))
        crop = crops[index] if index < len(crops) else None
        if crop:
            w, h = picture.size
            left, top, right, bottom = crop
            picture = picture.crop((round(w * left), round(h * top),
                                    round(w * (1 - right)), round(h * (1 - bottom))))
        pdf.drawImage(ImageReader(picture), x0, -y1, width=x1 - x0, height=y1 - y0, mask="auto")
