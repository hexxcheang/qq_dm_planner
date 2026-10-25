"""DM Optical document generator.

Upload a customer PO -> check the extracted details -> pick the date ->
download the Invoice and Delivery Order.

Run with:  streamlit run app.py
"""

from __future__ import annotations

import hashlib
from datetime import date
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import load_workbook

from docgen import VERSION
from docgen.generator import INVOICE_TEMPLATE, find_soffice, generate_documents
from docgen.models import Customer, LineItem, PurchaseOrder
from docgen.numbering import Register, do_no, invoice_no
from docgen.po_parser import SUPPORTED_TYPES, ParseError, claude_available, parse_po

ROOT = Path(__file__).resolve().parent
OUTPUT_DIR = ROOT / "output"
REGISTER = Register(ROOT / "data" / "register.json")
ITEM_COLUMNS = ["Description", "Qty", "Units", "Unit Price"]


@st.cache_data
def template_customer() -> Customer:
    ws = load_workbook(INVOICE_TEMPLATE).active
    return Customer(*[str(ws[f"A{r}"].value or "") for r in range(12, 16)])


def items_frame(po: PurchaseOrder) -> pd.DataFrame:
    rows = [[i.description, i.qty, i.units, i.unit_price] for i in po.items]
    return pd.DataFrame(rows or [["", 1.0, "Pcs", 0.0]], columns=ITEM_COLUMNS)


def frame_items(df: pd.DataFrame) -> list[LineItem]:
    items = []
    for row in df.itertuples(index=False):
        desc = str(row[0] or "").strip()
        if not desc:
            continue
        items.append(LineItem(description=desc, qty=float(row[1] or 0),
                              units=str(row[2] or "Pcs").strip() or "Pcs",
                              unit_price=float(row[3] or 0)))
    return items


st.set_page_config(page_title="DM Optical – Invoice & DO", page_icon="🧾", layout="wide")
st.title("Invoice & Delivery Order generator")

with st.sidebar:
    st.header("Settings")
    use_claude = st.toggle("Read PO with Claude AI (optional, paid)", value=claude_available(),
                           disabled=not claude_available(),
                           help="Needs ANTHROPIC_API_KEY. Not required: the free built-in reader "
                                "handles PDFs, scans, photos, Excel and CSV. Claude can help with "
                                "unusual PO layouts.")
    if not claude_available():
        st.caption("Using the free built-in reader (works offline, reads scans too).")
    has_soffice = find_soffice() is not None
    make_pdf = st.checkbox("Also create PDF files", value=has_soffice, disabled=not has_soffice,
                           help="Needs LibreOffice installed.")
    if not has_soffice:
        st.caption("Install LibreOffice to get PDFs as well as Excel files.")
    st.divider()
    st.caption(f"Version {VERSION}")

st.subheader("1. Upload the customer's Purchase Order")
upload = st.file_uploader("PO file", type=[t.lstrip(".") for t in SUPPORTED_TYPES])

if upload is None:
    st.info("Upload a PO (PDF, Excel, CSV or a photo) to begin.")
    st.stop()

data = upload.getvalue()
file_key = hashlib.sha1(data).hexdigest()
if st.session_state.get("file_key") != file_key or st.session_state.get("reader_claude") != use_claude:
    with st.spinner("Reading the PO… (scans take a few seconds per page)"):
        try:
            po, reader = parse_po(data, upload.name, use_claude=use_claude)
        except ParseError as exc:
            st.warning(str(exc))
            po, reader = PurchaseOrder(), "manual"
        except Exception as exc:  # network/API problems: let the user carry on by hand
            st.error(f"Could not read the PO automatically: {exc}")
            po, reader = PurchaseOrder(), "manual"
    st.session_state.update(file_key=file_key, reader_claude=use_claude, po=po, reader=reader,
                            generated=None)

po: PurchaseOrder = st.session_state["po"]
if st.session_state["reader"] != "manual":
    if po.items:
        st.success(f"Read {len(po.items)} item(s) from the PO ({st.session_state['reader']}). "
                   "Please check them below.")
    else:
        st.warning("No line items found – please type them in below.")

st.subheader("2. Check the details")
col1, col2, col3 = st.columns(3)
with col1:
    po_number = st.text_input("Customer PO No.", value=po.po_number, key=f"po_{file_key}")
with col2:
    doc_date = st.date_input("Invoice / DO date", value=date.today(), format="DD/MM/YYYY")
with col3:
    seq = st.number_input("Running number for this date", min_value=1, max_value=99, step=1,
                          value=REGISTER.next_seq(doc_date), key=f"seq_{doc_date}_{file_key}",
                          help="01 for the first PO on this date, 02 for the second, …")
st.markdown(f"**Invoice No.:** `{invoice_no(doc_date, seq)}` &nbsp;&nbsp; "
            f"**DO No.:** `{do_no(doc_date, seq)}`")
if seq <= REGISTER.last_seq(doc_date):
    st.warning(f"Number {seq:02d} has already been used for {doc_date:%d %b %Y}.")

edited = st.data_editor(
    items_frame(po), key=f"items_{file_key}", num_rows="dynamic", use_container_width=True,
    column_config={
        "Description": st.column_config.TextColumn(width="large", required=True),
        "Qty": st.column_config.NumberColumn(min_value=0, step=1, required=True),
        "Units": st.column_config.TextColumn(default="Pcs"),
        "Unit Price": st.column_config.NumberColumn(min_value=0, format="%.2f", required=True),
    },
)
items = frame_items(edited)
total = round(sum(i.amount for i in items), 2)
st.markdown(f"**Total / Amount Due:** ${total:,.2f}")

default_customer = template_customer()
with st.expander("Bill To / Ship To customer"):
    change_customer = st.checkbox(f"Use a different customer (default: {default_customer.name})")
    if change_customer:
        customer = Customer(
            name=st.text_input("Company", value=default_customer.name),
            address1=st.text_input("Address line 1", value=default_customer.address1),
            address2=st.text_input("Address line 2", value=default_customer.address2),
            attention=st.text_input("Attention", value=default_customer.attention),
        )
    else:
        customer = None

st.subheader("3. Generate")
problems = []
if not po_number.strip():
    problems.append("Enter the Customer PO No.")
if not items:
    problems.append("Add at least one line item.")
for p in problems:
    st.caption(f"⚠️ {p}")

if st.button("Generate Invoice & DO", type="primary", disabled=bool(problems)):
    final = PurchaseOrder(po_number=po_number.strip(), items=items, customer=customer)
    out_dir = OUTPUT_DIR / doc_date.isoformat()
    with st.spinner("Creating documents…"):
        docs = generate_documents(final, doc_date, int(seq), out_dir, make_pdf=make_pdf)
    REGISTER.record(doc_date, int(seq), final.po_number)
    st.session_state["generated"] = docs

docs = st.session_state.get("generated")
if docs:
    st.success(f"Created {docs.invoice_no} and {docs.do_no} (saved in {docs.invoice_xlsx.parent}).")
    xlsx_mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    c1, c2 = st.columns(2)
    for col, label, xlsx, pdf in ((c1, "Invoice", docs.invoice_xlsx, docs.invoice_pdf),
                                  (c2, "Delivery Order", docs.do_xlsx, docs.do_pdf)):
        with col:
            st.download_button(f"⬇️ {label} (Excel)", xlsx.read_bytes(), file_name=xlsx.name,
                               mime=xlsx_mime, use_container_width=True)
            if pdf:
                st.download_button(f"⬇️ {label} (PDF)", pdf.read_bytes(), file_name=pdf.name,
                                   mime="application/pdf", use_container_width=True)
