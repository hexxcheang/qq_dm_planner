"""Customer quotation: DMO price list lookup and quotation generation.

Uses small made-up workbooks shaped like the real DMO price list and
quotation template (those stay out of the repo).
"""

from datetime import date, datetime

import pdfplumber
import pytest
from openpyxl import Workbook
from openpyxl.styles import Border, Font, Side

from docgen.pricing import DmoCatalog, read_dmo_catalog, read_product_codes
from docgen.private_files import QUOTATION_TEMPLATE, classify
from docgen.quotation import (CustomerBook, Quotation, QuoteCustomer, QuoteLine, add_one_month,
                              build_quotation, generate_quotation)


@pytest.fixture
def dmo_list(tmp_path):
    wb = Workbook()
    fa = wb.active
    fa.title = "FA"
    fa["A1"], fa["E1"] = "Indicating the main models", "Indicating the instock models"
    fa["A2"] = "Indicating obsolete models"
    fa["A4"] = "FA LENS"
    fa.append(["Product Code (DM Optical)", "Model No. (Chiopt)", "Specifications", None, None,
               None, "JM Buying Price (SGD) ", "Remarks"])
    fa.append([None, None, "Focal length", "FNO", "Max Image Circle", "Resolution"])
    fa.append([' 1/1.8" 5MP FA Series'])
    fa.append([28160140, "FA0401C", "4mm", "F2.2", '1/1.8"', "5Mega", 47])
    fa.append(["28160140-1", "FA4501C", "4.5mm", "F2.2", '1/1.8"', "5Mega", 122])
    fa.append(['1.1" Macro Series'])
    fa.append([73620230, "ML5010C", "50mm", 2.8, '1.1"', "12M", 85])
    fa.append([None, "ML5010M1", None, None, '1.1"', "12M"])  # no code, no price
    acc = wb.create_sheet("Lens Accessories")
    acc["A1"] = "LENS ACCESSORIES"
    acc["A3"] = "Teleconverter"
    acc.append(["Model No. (Chiopt)", "Specifications", "JM Buying Price (SGD) ", "Remarks"])
    acc.append([None, "Length"])
    acc.append(["BA1501A", "1.5X", 40, "TELECONVERTER"])
    path = tmp_path / "dmo.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def quotation_template(tmp_path):
    wb = Workbook()
    ws = wb.active
    for col, width in zip("ABCDEFGH", [7.7, 9.2, 12.2, 9.5, 7.6, 10.4, 14.4, 12.8]):
        ws.column_dimensions[col].width = width  # as in the real template
    ws["G2"] = "DM Optical Pte Ltd"
    ws["A7"], ws["G7"] = "Attention", "Quotation"
    for row, text in zip(range(8, 13), ["Atrobot Sdn Bhd", "49G, Lebuh Bukit Kecil 2",
                                        "11900 Bayan Lepas", "Andrew Tan",
                                        "Andrew@atrobotvision.com"]):
        ws[f"A{row}"] = text
    ws["A12"].hyperlink = "mailto:Andrew@atrobotvision.com"
    ws["A12"].font = Font(underline="single", color="0563C1")
    ws["G9"], ws["H9"] = "Date", datetime(2026, 9, 23)
    ws["G10"], ws["H10"] = "Quotation No.", "Q260923-01"
    ws["G11"], ws["H11"] = "Valid Till", datetime(2026, 10, 22)
    ws["H9"].number_format = ws["H11"].number_format = "d-mmm-yy"
    for col, title in zip("ABEFGH", ["S/N", "Description", "Qty", "Units", "Unit Price", "Amount"]):
        ws[f"{col}14"] = title
    thin = Side(style="thin")
    for row in range(15, 19):
        for col in "ABEFGH":
            ws[f"{col}{row}"].border = Border(left=thin, right=thin,
                                              bottom=thin if row == 18 else None)
    ws["G19"], ws["H19"] = "Total (SGD)", "=SUM(H15:H18)"
    ws["H19"].number_format = '"$"#,##0.00_);[Red]\\("$"#,##0.00\\)'
    ws["A22"] = "Delivery Terms: within 2-3 weeks"
    ws.merge_cells("A22:G23")
    path = tmp_path / "quotation_template.xlsx"
    wb.save(path)
    return path


def test_catalog_reads_series_codes_and_prices(dmo_list):
    items = read_dmo_catalog(dmo_list)
    assert [(i.code, i.model, i.jm_price_sgd, i.series) for i in items] == [
        ("28160140", "FA0401C", 47.0, '1/1.8" 5MP FA Series'),
        ("28160140-1", "FA4501C", 122.0, '1/1.8" 5MP FA Series'),
        ("73620230", "ML5010C", 85.0, '1.1" Macro Series'),
        (None, "ML5010M1", None, '1.1" Macro Series'),
        (None, "BA1501A", 40.0, "Teleconverter"),
    ]
    assert read_product_codes(dmo_list) == {"28160140": "FA0401C", "28160140-1": "FA4501C",
                                            "73620230": "ML5010C"}


def test_quote_price_and_description(dmo_list):
    catalog = DmoCatalog.from_file(dmo_list)
    by_code = catalog.resolve("28160140")
    assert by_code.description == '28160140, 1/1.8" 5MP FA Series'
    assert by_code.unit_price == 72.31  # 47 / 0.65
    by_model = catalog.resolve("fa0401c")
    assert (by_model.description, by_model.unit_price, by_model.matched_by) == (
        '28160140, 1/1.8" 5MP FA Series', 72.31, "Chiopt model no.")
    assert catalog.resolve(73620230).unit_price == 130.77  # 85 / 0.65


def test_quote_without_code_or_price(dmo_list):
    catalog = DmoCatalog.from_file(dmo_list)
    accessory = catalog.resolve("BA1501A")
    assert accessory.description == "BA1501A, Teleconverter" and accessory.unit_price == 61.54
    assert "model no. used" in accessory.note
    no_price = catalog.resolve("ML5010M1")
    assert no_price.unit_price is None and "enter the price" in no_price.note
    assert catalog.resolve("XYZ").item is None


def test_valid_till_is_one_month_later():
    assert add_one_month(date(2026, 10, 25)) == date(2026, 11, 25)
    assert add_one_month(date(2026, 12, 15)) == date(2027, 1, 15)
    assert add_one_month(date(2026, 1, 31)) == date(2026, 2, 28)
    assert add_one_month(date(2028, 1, 31)) == date(2028, 2, 29)


def _quote() -> Quotation:
    customer = QuoteCustomer("JM Vistec System Pte Ltd", "10 Kaki Bukit Ave 1, #07-06",
                             "Singapore 417942", "Ms. Nicole Chan", "nicole@jmvistec.com")
    lines = [QuoteLine('28160140, 1/1.8" 5MP FA Series', 2, 72.31),
             QuoteLine('73620230, 1.1" Macro Series', 1, 130.77),
             QuoteLine("BA1501A, Teleconverter", 5, 61.54),
             QuoteLine('28160140-1, 1/1.8" 5MP FA Series', 3, 187.69)]
    return Quotation(date(2026, 10, 25), 1, customer, lines)


def test_build_quotation(quotation_template):
    q = _quote()
    ws = build_quotation(q, quotation_template).active
    assert [ws[f"A{r}"].value for r in range(8, 13)] == [
        "JM Vistec System Pte Ltd", "10 Kaki Bukit Ave 1, #07-06", "Singapore 417942",
        "Ms. Nicole Chan", "nicole@jmvistec.com"]
    assert ws["A12"].hyperlink.target == "mailto:nicole@jmvistec.com"
    assert ws["H9"].value == datetime(2026, 10, 25)
    assert ws["H10"].value == "Q261025-01"
    assert ws["H11"].value == datetime(2026, 11, 25)
    for i, line in enumerate(q.lines):
        r = 15 + i
        assert (ws[f"A{r}"].value, ws[f"B{r}"].value, ws[f"E{r}"].value, ws[f"F{r}"].value,
                ws[f"G{r}"].value, ws[f"H{r}"].value) == (
            i + 1, line.description, line.qty, "Pcs", line.unit_price, f"=G{r}*E{r}")
    total_row = next(r for r in range(15, ws.max_row + 1) if ws[f"G{r}"].value == "Total (SGD)")
    assert ws[f"H{total_row}"].value == "=SUM(H15:H18)"


def test_generate_quotation_pdf(tmp_path, quotation_template):
    pdf_path = generate_quotation(_quote(), quotation_template, tmp_path)
    assert pdf_path.name == "Q261025-01_JM_Vistec_System_Pte_Ltd.pdf"
    with pdfplumber.open(pdf_path) as pdf:
        text = pdf.pages[0].extract_text()
    for expected in ("Q261025-01", "25-Oct-26", "25-Nov-26", "JM Vistec System Pte Ltd",
                     "Ms. Nicole Chan", "nicole@jmvistec.com", '28160140, 1/1.8" 5MP FA Series',
                     "72.31", "144.62", "307.70", "$1,146.16"):
        assert expected in text, expected
    assert "Atrobot" not in text


def test_classify_quotation_template(quotation_template):
    assert classify(quotation_template) == QUOTATION_TEMPLATE


def test_customer_book_remembers_latest_first(tmp_path):
    book = CustomerBook(tmp_path / "customers.json")
    book.remember(QuoteCustomer("A Pte Ltd", email="a@a.com"))
    book.remember(QuoteCustomer("B Pte Ltd"))
    book.remember(QuoteCustomer("a pte ltd", email="new@a.com"))
    assert [(c.company, c.email) for c in book.all()] == [("a pte ltd", "new@a.com"),
                                                         ("B Pte Ltd", "")]
