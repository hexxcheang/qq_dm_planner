from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from docgen.generator import generate_documents
from docgen.models import Customer, LineItem, PurchaseOrder
from docgen.numbering import Register, do_no, invoice_no
from docgen.po_parser import ParseError, parse_with_rules

from . import sample_po

DAY = date(2026, 9, 29)


def test_document_numbers():
    assert invoice_no(DAY, 1) == "IN260929-01"
    assert do_no(DAY, 12) == "DO260929-12"


def test_register_increments_per_date(tmp_path):
    reg = Register(tmp_path / "register.json")
    assert reg.next_seq(DAY) == 1
    reg.record(DAY, 1, "PO-A")
    reg.record(DAY, 2, "PO-B")
    assert reg.next_seq(DAY) == 3
    assert reg.next_seq(date(2026, 9, 30)) == 1


def _po(n_items: int) -> PurchaseOrder:
    return PurchaseOrder("POR13413", [LineItem(f"Part {i}", i, "Pcs", 10.0 * i)
                                      for i in range(1, n_items + 1)])


@pytest.mark.parametrize("n_items", [1, 2, 3, 8])
def test_invoice_fields(tmp_path, n_items):
    docs = generate_documents(_po(n_items), DAY, 3, tmp_path)
    ws = load_workbook(docs.invoice_xlsx).active
    assert ws["C5"].value == "POR13413"
    assert ws["H10"].value == datetime(2026, 9, 29)
    assert ws["H11"].value == "IN260929-03"
    assert ws["H12"].value == "DO260929-03"
    for i in range(n_items):
        row = 18 + i
        assert ws[f"A{row}"].value == i + 1
        assert ws[f"B{row}"].value == f"Part {i + 1}"
        assert ws[f"E{row}"].value == i + 1
        assert ws[f"F{row}"].value == "Pcs"
        assert ws[f"G{row}"].value == 10 * (i + 1)
        assert ws[f"H{row}"].value == f"=G{row}*E{row}"
        assert f"B{row}:D{row}" in {r.coord for r in ws.merged_cells.ranges}
    total_row = next(r for r in range(18, ws.max_row + 1) if ws[f"G{r}"].value == "Total")
    assert ws[f"H{total_row}"].value == f"=SUM(H18:H{18 + n_items - 1})"
    assert ws[f"G{total_row + 1}"].value == "Amount Due"
    assert ws[f"H{total_row + 1}"].value == f"=H{total_row}"
    # Terms block still below the totals, still merged.
    terms = [r for r in ws.merged_cells.ranges if r.min_col == 1 and r.max_col == 7]
    assert terms and terms[0].min_row == total_row + 3


@pytest.mark.parametrize("n_items", [1, 5])
def test_delivery_order_fields(tmp_path, n_items):
    docs = generate_documents(_po(n_items), DAY, 1, tmp_path)
    ws = load_workbook(docs.do_xlsx).active
    assert ws["G9"].value == datetime(2026, 9, 29)
    assert ws["G10"].value == "DO260929-01"
    assert ws["G11"].value == "POR13413"
    for i in range(n_items):
        row = 16 + i
        assert ws[f"A{row}"].value == i + 1
        assert ws[f"B{row}"].value == i + 1
        assert ws[f"C{row}"].value == "Pcs"
        assert ws[f"D{row}"].value == f"Part {i + 1}"
        assert ws[f"F{row}"].value == "N/A"
    assert ws[f"A{16 + n_items}"].value is None
    received = next(r for r in range(16, ws.max_row + 1)
                    if ws[f"F{r}"].value == "RECEIVED IN GOOD ORDER AND CONDITION")
    assert ws[f"A{received}"].value == "Ship Via"


def test_customer_override(tmp_path):
    po = _po(1)
    po.customer = Customer("ACME PTE LTD", "1 Main St", "Singapore 123456", "Jane")
    docs = generate_documents(po, DAY, 1, tmp_path)
    inv = load_workbook(docs.invoice_xlsx).active
    do = load_workbook(docs.do_xlsx).active
    assert [inv[f"A{r}"].value for r in range(12, 16)] == po.customer.lines()
    assert [inv[f"D{r}"].value for r in range(12, 16)] == po.customer.lines()
    assert [do[f"A{r}"].value for r in range(10, 14)] == po.customer.lines()


def _check_sample(po: PurchaseOrder, with_units: bool = True):
    assert po.po_number == sample_po.PO_NUMBER
    assert [(i.description, i.qty, i.unit_price) for i in po.items] == [
        (d, q, p) for d, q, _u, p in sample_po.ITEMS]
    if with_units:
        assert [i.units for i in po.items] == [u for _d, _q, u, _p in sample_po.ITEMS]


def test_rules_parser_pdf_ruled():
    _check_sample(parse_with_rules(sample_po.make_pdf(ruled=True), "po.pdf"))


def test_rules_parser_pdf_unruled():
    _check_sample(parse_with_rules(sample_po.make_pdf(ruled=False), "po.pdf"))


def test_rules_parser_xlsx():
    _check_sample(parse_with_rules(sample_po.make_xlsx(), "po.xlsx"))


def test_rules_parser_csv():
    po = parse_with_rules(sample_po.make_csv(), "po.csv")
    _check_sample(po, with_units=False)
    assert all(i.units == "Pcs" for i in po.items)


def test_rules_parser_rejects_images():
    with pytest.raises(ParseError):
        parse_with_rules(b"\x89PNG", "po.png")
