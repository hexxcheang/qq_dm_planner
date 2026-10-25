"""Supplier PO: price lookup and PO generation.

Uses small made-up workbooks shaped like the real price lists and PO template
(those hold real prices and stay out of the repo).
"""

from datetime import date, datetime

import pdfplumber
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Side

from docgen.private_files import (CHIOPT_PRICES, DMO_PRICES, PO_TEMPLATE, QUOTATION_TEMPLATE,
                                  PrivateStore, classify)
from docgen.pricing import PriceBook
from docgen.supplier_po import POLine, SupplierPO, build_supplier_po, generate_supplier_po


@pytest.fixture
def chiopt_list(tmp_path):
    wb = Workbook()
    fa = wb.active
    fa.title = "FA"
    fa["G2"], fa["H2"] = "1SGD=", 5.5
    fa.append([])
    fa["A5"] = ' 1/2" FA系列'
    fa.append(["型号（Model No.）", "规格（Spec）", None, None, None, "市场指导价", "国内含税\n售价",
               "向新加坡子公司新币售价", "备注", "DM Buying Price (USD)"])
    fa.append([None, "Focal length", "FNO"])
    fa.append(["FA0401C", "4mm", "F2.2", None, None, None, 210, 37.2, None, 29.016])
    fa.append(["1.1\" ML系列"])
    fa.append(["型号（Model No.）", "规格（Spec）"])  # later table: price column not repeated
    fa.append(["ML5010C", "50mm", 2.8, None, None, None, 380, 67.3, None, 52.494])
    micro = wb.create_sheet("显微镜")
    micro.append(["型号（Model No.）", "规格（Spec）", "国内含税售价"])
    micro.append(["CA10007A", "10X", 5000])
    path = tmp_path / "chiopt.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def dmo_list(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "FA"
    ws["A1"] = "Indicating the main models"
    ws.append([])
    ws.append(["FA LENS"])
    ws.append(["Product Code (DM Optical)", "Model No. (Chiopt)", "Specifications", None, None,
               None, "JM Buying Price (SGD) "])
    ws.append([28160140, "FA0401C", "4mm"])
    ws.append([73620230, "ML5010C", "50mm"])
    ws.append(["80770340-1", "MV65021B"])
    ws.append(["-", "G-YT0802C"])
    path = tmp_path / "dmo.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def po_template(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["D1"] = "DM Optical Pte Ltd"
    ws["E5"] = "Purchase Order"
    ws["A7"] = "Hunan Chiopt Optical Technology Co., Ltd"
    ws["A8"] = "湖南省长沙市雨花区"
    ws["E6"], ws["G6"] = "PO#", "PO261009-01"
    ws["E7"], ws["G7"] = "Date", datetime(2026, 10, 9)
    ws["E8"], ws["G8"] = "Delivery Date", datetime(2026, 10, 19)
    ws["G7"].number_format = ws["G8"].number_format = "d-mmm-yy"
    for col, title in zip("ABDEFH", ["S/N", "Description", "Quantity", "UOM", "Unit Price",
                                     "Amount"]):
        ws[f"{col}15"] = title
    thin = Side(style="thin")
    for row in range(16, 20):
        for col in "ADEFGH":
            ws[f"{col}{row}"].border = Border(left=thin, right=thin,
                                              bottom=thin if row == 19 else None)
    ws["A16"], ws["B16"], ws["D16"], ws["E16"], ws["G16"] = 1, "ML5010C", 1, "Pcs", 52.49
    ws["H16"] = "=D16*G16"
    ws["A19"] = "REMARK :\nSHIPPING ADDRESS :\nJM VISTEC SYSTEM PTE LTD"
    ws.merge_cells("A19:C19")
    ws["G20"], ws["H20"] = "Total (USD)", "=SUM(H16:H19)"
    path = tmp_path / "po_template.xlsx"
    wb.save(path)
    return path


def test_lookup_by_product_code_and_model(chiopt_list, dmo_list):
    book = PriceBook.from_files(chiopt_list, dmo_list)

    by_code = book.resolve("73620230")
    assert (by_code.model, by_code.price_usd, by_code.matched_by) == (
        "ML5010C", 52.49, "DM product code")
    assert book.resolve(73620230).model == "ML5010C"  # pasted from Excel as a number

    by_model = book.resolve(" ml5010c ")
    assert (by_model.model, by_model.price_usd, by_model.matched_by) == (
        "ML5010C", 52.49, "Chiopt model no.")

    assert book.resolve("28160140").price_usd == 29.02
    assert book.resolve("FA0401C").price_usd == 29.02


def test_lookup_without_price(chiopt_list, dmo_list):
    book = PriceBook.from_files(chiopt_list, dmo_list)
    no_price = book.resolve("CA10007A")  # sheet without a DM Buying Price column
    assert no_price.model == "CA10007A" and no_price.price_usd is None and no_price.note
    not_in_chiopt = book.resolve("80770340-1")
    assert not_in_chiopt.model == "MV65021B" and not_in_chiopt.price_usd is None
    unknown = book.resolve("XYZ123")
    assert unknown.model is None and "Not found" in unknown.note
    assert book.resolve("-").model is None


def test_classify_and_store(tmp_path, chiopt_list, dmo_list, po_template):
    assert classify(chiopt_list) == CHIOPT_PRICES
    assert classify(dmo_list) == DMO_PRICES
    assert classify(po_template) == PO_TEMPLATE
    store = PrivateStore(tmp_path / "private")
    supplier_files = (PO_TEMPLATE, CHIOPT_PRICES, DMO_PRICES)
    assert store.missing(supplier_files) == list(supplier_files)
    for path in (chiopt_list, dmo_list, po_template):
        store.save(path, path.name)
    assert store.missing(supplier_files) == []
    assert store.missing() == [QUOTATION_TEMPLATE]
    assert store.describe(PO_TEMPLATE).startswith("po_template.xlsx, saved")


def test_build_supplier_po(po_template):
    lines = [POLine("ML5010C", 2, 52.49), POLine("FA0401C", 10, 29.02), POLine("CA10007A", 3, 120),
             POLine("HC1205A", 1, 57.95)]
    po = SupplierPO(date(2026, 10, 25), 1, lines)
    ws = build_supplier_po(po, po_template).active
    assert ws["G6"].value == "PO261025-01"
    assert ws["G7"].value == datetime(2026, 10, 25)
    assert ws["G8"].value == datetime(2026, 11, 4)  # +10 days
    for i, line in enumerate(lines):
        row = 16 + i
        assert (ws[f"A{row}"].value, ws[f"B{row}"].value, ws[f"D{row}"].value,
                ws[f"E{row}"].value, ws[f"G{row}"].value) == (i + 1, line.model, line.qty, "Pcs",
                                                               line.unit_price)
        assert ws[f"H{row}"].value == f"=D{row}*G{row}"
    total_row = next(r for r in range(16, ws.max_row + 1) if ws[f"G{r}"].value == "Total (USD)")
    assert ws[f"H{total_row}"].value == "=SUM(H16:H19)"
    remark = next(r for r in ws.merged_cells.ranges if r.min_col == 1 and r.max_col == 3)
    assert ws.cell(remark.min_row, 1).value.startswith("REMARK")
    assert remark.min_row == total_row - 1


def test_generate_supplier_po_pdf(tmp_path, po_template):
    po = SupplierPO(date(2026, 10, 25), 2, [POLine("ML5010C", 2, 52.49), POLine("FA0401C", 10, 29.02)])
    pdf_path = generate_supplier_po(po, po_template, tmp_path)
    assert pdf_path.name == "PO261025-02_Chiopt.pdf"
    with pdfplumber.open(pdf_path) as pdf:
        text = pdf.pages[0].extract_text()
    for expected in ("PO261025-02", "25-Oct-26", "4-Nov-26", "ML5010C", "FA0401C", "52.49",
                     "104.98", "290.20", "395.18", "湖南省长沙市雨花区"):
        assert expected in text, expected
    assert load_workbook(po_template).active["G6"].value == "PO261009-01"  # template untouched
